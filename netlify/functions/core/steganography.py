"""
Simplified Steganography Module for Netlify Functions
"""

import os
import json
import base64
from datetime import datetime
from PIL import Image
import numpy as np
import openpyxl
import csv
import PyPDF2


class DocumentSteganography:
    """Simplified steganography for serverless environment"""
    
    def __init__(self):
        self.delimiter = b'\x00\x00\xFF\xFF\x00\x00'
    
    def hide_data_in_image(self, cover_path: str, secret_data: bytes, output_path: str, original_hash: str = None) -> dict:
        """Hide data in image using LSB steganography"""
        try:
            image = Image.open(cover_path)
            if image.mode != 'RGB':
                image = image.convert('RGB')
            
            img_array = np.array(image)
            height, width, channels = img_array.shape
            
            # Prepare metadata
            metadata = {
                'original_name': os.path.basename(cover_path),
                'data_length': len(secret_data),
                'hidden_at': datetime.now().isoformat(),
                'method': 'lsb',
                'original_hash': original_hash
            }
            
            metadata_json = json.dumps(metadata).encode('utf-8')
            full_data = metadata_json + self.delimiter + secret_data
            
            # Check capacity
            max_capacity = (height * width * channels) // 8
            if len(full_data) > max_capacity:
                return {
                    'success': False,
                    'error': f'Data too large. Max: {max_capacity}, Data: {len(full_data)}'
                }
            
            # Hide data using LSB
            binary_data = ''.join(format(byte, '08b') for byte in full_data)
            img_flat = img_array.flatten()
            
            for i, bit in enumerate(binary_data):
                img_flat[i] = (img_flat[i] & 0xFE) | int(bit)
            
            stego_array = img_flat.reshape(height, width, channels)
            stego_image = Image.fromarray(stego_array.astype('uint8'))
            
            # Save as PNG to preserve data
            output_path = os.path.splitext(output_path)[0] + '.png'
            stego_image.save(output_path, 'PNG')
            
            return {
                'success': True,
                'stego_image': output_path,
                'method': 'lsb',
                'data_size': len(secret_data)
            }
            
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def extract_data_from_image(self, stego_path: str) -> dict:
        """Extract data from steganographic image"""
        try:
            image = Image.open(stego_path)
            if image.mode != 'RGB':
                image = image.convert('RGB')
            
            img_array = np.array(image)
            img_flat = img_array.flatten()
            
            # Extract LSBs
            binary_data = ''.join(str(pixel & 1) for pixel in img_flat)
            
            # Convert to bytes
            extracted_bytes = bytearray()
            for i in range(0, len(binary_data), 8):
                byte_chunk = binary_data[i:i+8]
                if len(byte_chunk) == 8:
                    extracted_bytes.append(int(byte_chunk, 2))
            
            # Find delimiter
            delimiter_index = extracted_bytes.find(self.delimiter)
            if delimiter_index == -1:
                return {'success': False, 'error': 'No hidden data found'}
            
            # Extract metadata and secret data
            metadata_bytes = extracted_bytes[:delimiter_index]
            metadata = json.loads(metadata_bytes.decode('utf-8'))
            
            secret_start = delimiter_index + len(self.delimiter)
            secret_end = secret_start + metadata['data_length']
            secret_data = bytes(extracted_bytes[secret_start:secret_end])
            
            return {
                'success': True,
                'secret_data': secret_data,
                'metadata': metadata,
                'method': 'lsb'
            }
            
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def hide_data_in_pdf(self, cover_path: str, secret_data: bytes, output_path: str, original_hash: str = None) -> dict:
        """Hide data in PDF metadata"""
        try:
            with open(cover_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                writer = PyPDF2.PdfWriter()
                
                # Copy pages
                for page in reader.pages:
                    writer.add_page(page)
                
                # Prepare metadata
                metadata = {
                    'original_name': os.path.basename(cover_path),
                    'data_length': len(secret_data),
                    'hidden_at': datetime.now().isoformat(),
                    'method': 'pdf_metadata',
                    'original_hash': original_hash
                }
                
                metadata_json = json.dumps(metadata).encode('utf-8')
                full_data = metadata_json + self.delimiter + secret_data
                encoded_data = base64.b64encode(full_data).decode('utf-8')
                
                # Add to metadata
                writer.add_metadata({
                    '/Subject': encoded_data,
                    '/Keywords': 'document_integrity_protected'
                })
                
                with open(output_path, 'wb') as output_file:
                    writer.write(output_file)
            
            return {
                'success': True,
                'stego_document': output_path,
                'method': 'pdf_metadata'
            }
            
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def extract_data_from_pdf(self, stego_path: str) -> dict:
        """Extract data from PDF metadata"""
        try:
            with open(stego_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                metadata = reader.metadata
                
                if not metadata or '/Subject' not in metadata:
                    return {'success': False, 'error': 'No hidden data found'}
                
                encoded_data = metadata['/Subject']
                full_data = base64.b64decode(encoded_data.encode('utf-8'))
                
                # Find delimiter
                delimiter_index = full_data.find(self.delimiter)
                if delimiter_index == -1:
                    return {'success': False, 'error': 'Data corrupted'}
                
                # Extract metadata and secret data
                metadata_bytes = full_data[:delimiter_index]
                doc_metadata = json.loads(metadata_bytes.decode('utf-8'))
                
                secret_start = delimiter_index + len(self.delimiter)
                secret_end = secret_start + doc_metadata['data_length']
                secret_data = full_data[secret_start:secret_end]
                
                return {
                    'success': True,
                    'secret_data': secret_data,
                    'metadata': doc_metadata,
                    'method': 'pdf_metadata'
                }
                
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def hide_data_in_excel(self, cover_path: str, secret_data: bytes, output_path: str, original_hash: str = None) -> dict:
        """Hide data in Excel file"""
        try:
            # Simple implementation - add data as hidden sheet
            wb = openpyxl.load_workbook(cover_path)
            
            # Create hidden sheet
            hidden_sheet = wb.create_sheet("__hidden__")
            hidden_sheet.sheet_state = 'hidden'
            
            # Prepare metadata
            metadata = {
                'original_name': os.path.basename(cover_path),
                'data_length': len(secret_data),
                'hidden_at': datetime.now().isoformat(),
                'method': 'excel_hidden_sheet',
                'original_hash': original_hash
            }
            
            metadata_json = json.dumps(metadata).encode('utf-8')
            full_data = metadata_json + self.delimiter + secret_data
            encoded_data = base64.b64encode(full_data).decode('utf-8')
            
            # Store in hidden cell
            hidden_sheet['A1'] = encoded_data
            
            wb.save(output_path)
            
            return {
                'success': True,
                'stego_excel': output_path,
                'method': 'excel_hidden_sheet'
            }
            
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def extract_data_from_excel(self, stego_path: str) -> dict:
        """Extract data from Excel file"""
        try:
            wb = openpyxl.load_workbook(stego_path)
            
            if "__hidden__" not in wb.sheetnames:
                return {'success': False, 'error': 'No hidden data found'}
            
            hidden_sheet = wb["__hidden__"]
            encoded_data = hidden_sheet['A1'].value
            
            if not encoded_data:
                return {'success': False, 'error': 'No hidden data found'}
            
            full_data = base64.b64decode(encoded_data.encode('utf-8'))
            
            # Find delimiter
            delimiter_index = full_data.find(self.delimiter)
            if delimiter_index == -1:
                return {'success': False, 'error': 'Data corrupted'}
            
            # Extract metadata and secret data
            metadata_bytes = full_data[:delimiter_index]
            metadata = json.loads(metadata_bytes.decode('utf-8'))
            
            secret_start = delimiter_index + len(self.delimiter)
            secret_end = secret_start + metadata['data_length']
            secret_data = full_data[secret_start:secret_end]
            
            return {
                'success': True,
                'secret_data': secret_data,
                'metadata': metadata,
                'method': 'excel_hidden_sheet'
            }
            
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def hide_data_in_docx(self, cover_path: str, secret_data: bytes, output_path: str, original_hash: str = None) -> dict:
        """Placeholder for DOCX steganography"""
        return {'success': False, 'error': 'DOCX steganography not implemented in serverless version'}
    
    def extract_data_from_docx(self, stego_path: str) -> dict:
        """Placeholder for DOCX extraction"""
        return {'success': False, 'error': 'DOCX extraction not implemented in serverless version'}